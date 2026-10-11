"""Non-PDF document extractors.

Each extractor returns list[tuple[int, str]] (page_number, text),
matching the contract of pdf_utils._extract_pymupdf.
Non-paginated formats return [(1, full_text)].
"""

import base64
import re
import shutil
import subprocess
import sys
import tempfile
from html.parser import HTMLParser
from pathlib import Path

from pdf_utils import clean_text

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".docx",
    ".pptx",
    ".rtf",
    ".odt",
    ".htm",
    ".html",
    ".epub",
    ".doc",
}


def extract_text(file_path: Path) -> list[tuple[int, str]]:
    """Dispatch extraction by file extension."""
    ext = file_path.suffix.lower()
    extractors = {
        ".txt": _extract_txt,
        ".md": _extract_txt,
        ".docx": _extract_docx,
        ".pptx": _extract_pptx,
        ".rtf": _extract_rtf,
        ".odt": _extract_odt,
        ".htm": _extract_html,
        ".html": _extract_html,
        ".epub": _extract_epub,
        ".doc": _extract_doc,
    }
    extractor = extractors.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported extension: {ext!r}")
    pages = extractor(file_path)
    # Validation: warn on suspiciously small output
    total_text = " ".join(t for _, t in pages)
    file_size = file_path.stat().st_size
    if file_size > 1024 and len(total_text.strip()) < 50:
        print(
            f"  WARNING: extracted only {len(total_text)} chars from {file_path.name} ({file_size} bytes)",
            file=sys.stderr,
        )
    return pages


def _extract_epub(file_path: Path) -> list[tuple[int, str]]:
    """Extract the EPUB reading order with the locally installed Pandoc."""
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise RuntimeError("EPUB extraction requires pandoc on PATH")
    result = subprocess.run(
        [pandoc, "--from=epub", "--to=plain", "--wrap=none", str(file_path.resolve())],
        capture_output=True,
        encoding="utf-8",
        check=True,
        timeout=180,
    )
    text = clean_text(result.stdout.strip())
    return [(1, text)] if text else []


def _extract_doc(file_path: Path) -> list[tuple[int, str]]:
    """Read legacy Word text (including tables) from a disposable source copy.

    Word is isolated from any interactive instance, opens read-only with macros
    disabled, and closes without saving. The source cannot acquire lock files.
    """
    if sys.platform != "win32":
        raise RuntimeError("DOC extraction requires Microsoft Word on Windows")
    helper = Path(__file__).with_name("extract_word.ps1")
    with tempfile.TemporaryDirectory(prefix="neph-word-") as tmp:
        copy = Path(tmp) / "source.doc"
        shutil.copyfile(file_path, copy)
        result = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(helper),
                "-DocumentPath",
                str(copy),
            ],
            capture_output=True,
            check=True,
            timeout=180,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    text = base64.b64decode(result.stdout.strip(), validate=True).decode("utf-8")
    text = clean_text(text.replace("\r", "\n").replace("\x07", " | ").strip())
    return [(1, text)] if text else []


def _read_text_file(file_path: Path) -> str:
    """Read a text file with UTF-8, falling back to latin-1."""
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1")


def _extract_txt(file_path: Path) -> list[tuple[int, str]]:
    """Extract text from plain text or markdown files."""
    text = clean_text(_read_text_file(file_path).strip())
    if not text:
        return []
    return [(1, text)]


def _extract_docx(file_path: Path) -> list[tuple[int, str]]:
    """Extract text from DOCX using python-docx (paragraphs + tables)."""
    from docx import Document

    doc = Document(str(file_path))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            cells = []
            for cell in row.cells:
                if id(cell) not in seen and cell.text.strip():
                    seen.add(id(cell))
                    cells.append(cell.text.strip())
            if cells:
                parts.append(" | ".join(cells))
    text = clean_text("\n".join(parts).strip())
    if not text:
        return []
    return [(1, text)]


def _extract_pptx(file_path: Path) -> list[tuple[int, str]]:
    """Extract text from PPTX using python-pptx (all slides)."""
    from pptx import Presentation

    prs = Presentation(str(file_path))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        parts.append(para.text)
    text = clean_text("\n".join(parts).strip())
    if not text:
        return []
    return [(1, text)]


def _extract_rtf(file_path: Path) -> list[tuple[int, str]]:
    """Extract text from RTF using striprtf."""
    from striprtf.striprtf import rtf_to_text

    raw = _read_text_file(file_path)
    text = clean_text(rtf_to_text(raw).strip())
    if not text and "\\shptxt" in raw:
        # striprtf ignores drawing containers, including their text boxes.
        # Expose only the primary shape text; keep shprslt (legacy duplicate),
        # picture data and shape properties ignorable.
        shape_text = re.sub(r"\\\*\\shpinst\b|\\(?:shp|shptxt)\b", "", raw)
        text = clean_text(rtf_to_text(shape_text).strip())
    if not text:
        return []
    return [(1, text)]


def _extract_odt(file_path: Path) -> list[tuple[int, str]]:
    """Extract text from ODT using odfpy."""
    from odf.opendocument import load
    from odf.text import P

    doc = load(str(file_path))
    parts = []
    for elem in doc.getElementsByType(P):
        text = "".join(node.data for node in elem.childNodes if hasattr(node, "data"))
        if text.strip():
            parts.append(text)
    text = clean_text("\n".join(parts).strip())
    if not text:
        return []
    return [(1, text)]


class _HTMLTextExtractor(HTMLParser):
    """HTML parser that extracts text, skipping script/style tags."""

    def __init__(self):
        super().__init__()
        self._parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip_depth += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip_depth = max(0, self._skip_depth - 1)

    def handle_data(self, data):
        if self._skip_depth == 0 and data.strip():
            self._parts.append(data.strip())

    def get_text(self) -> str:
        return "\n".join(self._parts)


def _extract_html(file_path: Path) -> list[tuple[int, str]]:
    """Extract text from HTML, skipping script/style tags."""
    raw = _read_text_file(file_path)
    parser = _HTMLTextExtractor()
    parser.feed(raw)
    text = clean_text(parser.get_text().strip())
    if not text:
        return []
    return [(1, text)]
